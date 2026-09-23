# -*- coding: utf-8 -*-
"""Class 04 — Descriptive geometry (U03), first of two classes. P2a due; P3a issued."""
import os, math
from nb import *
from slidekit import *
import coursedata as C
C.require_current_deck(4)
import diagrams as G
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

OUT = os.environ.get("DECK_OUT", "/home/claude/out")
d = Deck()
WK = 4
VISUAL = ("Visual quality 35% · computational understanding 25% · technical execution 20% "
          "· experimentation 10% · requirements and identity 10%.")

# 1 ── TITLE
title_slide(d, WK, ["DESCRIPTIVE", "GEOMETRY"],
            "P3 BEGINS — WHAT A SOLID IS MADE OF, UNDERNEATH", bg=YELLOW, lfill=CYAN, rfill=LIME)

# 2 ── THE LADDER
s = d.slide(BLACK)
s.header("POINT → LINE → EDGE → FACE → SOLID → BOOLEAN",
         "Six steps. Each one is built from the step before it.", tsize=30)
rungs = [("POINT", CYAN, "three numbers", "[x, y, z]"),
         ("LINE", LIME, "two positions", "[p1, p2]"),
         ("EDGE", YELLOW, "two vertices", "[v1, v2]"),
         ("FACE", PINK, "ordered corners", "[v1, v2, v3]"),
         ("SOLID", CYAN, "closed faces", "[f1, f2, f3…]"),
         ("BOOLEAN", LIME, "two solids", "A − B")]
x = L
wd = (W - 0.2 * 5) / 6
for nm, col, sub, code in rungs:
    s.chip(x, 2.1, wd, 0.5, nm, fill=col, size=12.5)
    s.card(x, 2.7, wd, 1.9, fill=CREAM)
    s.t(x + 0.15, 2.95, wd - 0.3, 0.5, [s.Cb(sub, 12.5)], align=PP_ALIGN.CENTER)
    s.t(x + 0.15, 3.6, wd - 0.3, 0.5, [s.Mb(code, 11.5)], align=PP_ALIGN.CENTER)
    x += wd + 0.2
s.banner(4.95, "A POSITION IS NOT A DISPLACEMENT. THE SAME THREE NUMBERS MEAN DIFFERENT THINGS "
              "DEPENDING ON WHICH ONE YOU MEANT.", fill=YELLOW, h=0.68, align=PP_ALIGN.CENTER)
s.t(L, 5.9, W, 0.9, [s.C("Nothing here is new geometry. It is the same collection, read at a "
                         "different level: a face is a list of vertices, a solid is a closed list "
                         "of faces, and in Week 6 an array will be a list of copies.", 13.5, MUTE)], ls=1.3)

# 3 ── A FACE IS AN ORDERED LIST
s = d.slide(CREAM)
s.header("A FACE IS AN ORDERED LIST OF CORNERS",
         "Same three vertices, two different faces. The order is the difference.")
s.panel(L, 1.95, 6.0, 3.4, "WINDING ORDER", headfill=CYAN)
cx0, cy0 = L + 1.4, 3.75
tri = [(cx0, cy0 - 0.7), (cx0 - 0.75, cy0 + 0.55), (cx0 + 0.75, cy0 + 0.55)]
for i, (px, py) in enumerate(tri):
    s.dot(px, py, d=0.16, fill=BLACK)
    s.t(px - 0.35, py - 0.42, 0.7, 0.3, [s.Mb("v%d" % i, 11)], align=PP_ALIGN.CENTER)
s.t(L + 3.0, 2.75, 2.8, 2.3, [
    s.Mb("f = [v0, v1, v2]", 12, bold=True),
    s.Mb("  normal → toward you", 11),
    s.Mb(" ", 8),
    s.Mb("f = [v0, v2, v1]", 12, bold=True),
    s.Mb("  normal → away", 11)], ls=1.4)
s.panel(7.0, 1.95, 5.65, 3.4, "WHY IT MATTERS LATER", headfill=PINK, tsize=14.5)
s.t(7.3, 2.72, 5.05, 2.4, [
    s.Cb("A flipped face tells a renderer the inside is the outside.", 14),
    s.Cb(" ", 7),
    s.Cb("It tells a Boolean it cannot decide what to keep.", 14),
    s.Cb(" ", 7),
    s.Cb("Most broken renders and failed cuts this semester will start here.", 14)], ls=1.3)
s.banner(5.65, "CLOSED MEANS EVERY EDGE IS SHARED BY EXACTLY TWO FACES. CHECK IT BEFORE YOU BELIEVE IT.",
         fill=LIME, h=0.72, align=PP_ALIGN.CENTER)

# 4 ── THE PYRAMID
s = d.slide(BLACK)
s.header("THE PYRAMID", "Five vertices, eight edges, five faces. Built by hand, driven by three numbers.")
s.panel(L, 2.05, 6.2, 3.3, "WHAT YOU ASSEMBLE", headfill=LIME)
counts = [("5", "VERTICES", "four base corners and an apex"),
          ("8", "EDGES", "four around the base, four to the apex"),
          ("5", "FACES", "four triangular sides and the base")]
y = 2.85
for n, nm, body in counts:
    s.chip(L + 0.35, y, 0.75, 0.62, n, fill=CYAN, size=17)
    s.t(L + 1.35, y, 4.6, 0.3, [s.Ab(nm, 12.5)])
    s.t(L + 1.35, y + 0.3, 4.6, 0.3, [s.Cb(body, 11.5)])
    y += 0.78
s.panel(7.15, 2.05, 5.5, 3.3, "EXPOSED INPUTS", headfill=YELLOW)
s.t(7.45, 2.85, 4.9, 2.2, [
    s.Mb("Width", 14, bold=True), s.Mb("Depth", 14, bold=True), s.Mb("Height", 14, bold=True),
    s.Mb(" ", 8),
    s.Cb("Three numbers, and the whole solid rebuilds. That is the entire point of the exercise.", 13)], ls=1.4)
s.banner(5.7, "DO NOT FORGET THE BASE FACE. A PYRAMID WITH FOUR SIDES AND NO BOTTOM IS NOT A SOLID.",
         fill=PINK, h=0.7, align=PP_ALIGN.CENTER)

# 5 ── THE BOOLEAN
s = d.slide(CREAM)
s.header("THEN IT GOES INTO A BOOLEAN",
         "Cube minus pyramid. The same subtraction as in P2, with a solid you built point by point.")
flow = [("01", "BUILD IT", CYAN, "Five vertices, five faces, closed, with Width / Depth / Height exposed."),
        ("02", "SUBTRACT IT", LIME, "Cube minus pyramid. Name which solid is kept and which is removed."),
        ("03", "DRIVE IT", YELLOW, "Change the three inputs and show at least three settings of the result."),
        ("04", "EXPLAIN IT", PINK, "Face order, normals, closure — why it cuts cleanly instead of failing.")]
x = L
wd = (W - 0.23 * 3) / 4
for n, ttl, col, body in flow:
    s.chip(x, 1.95, wd, 0.5, n + "   " + ttl, fill=col, size=12.5)
    s.card(x, 2.55, wd, 2.2, fill=CREAM)
    s.t(x + 0.25, 2.55, wd - 0.5, 2.2, [s.Cb(body, 13)], ls=1.3, anchor=MSO_ANCHOR.MIDDLE)
    x += wd + 0.23
s.banner(5.05, "AN OPEN MESH HAS NO INSIDE. A BOOLEAN NEEDS ONE.",
         fill=BLACK, color=CYAN, h=0.68, align=PP_ALIGN.CENTER)
s.t(L, 5.95, W, 0.8, [s.C("The pyramid is the warm-up. Next week the same two lists describe your own "
                          "module — the one you light and render for P3a, and array for P3b.",
                          13.5, GREY)], ls=1.3)

# 6 ── THE SEQUENCE
unit_steps_slide(d, "U03", title="THE SEQUENCE", n=(0, 6),
                 sub="What we build together, in order, this class.")

# 7 ── PSEUDOCODE
pseudocode_slide(d, "U03",
                 sub="A mesh is two lists. That is the entire idea, and it fits on one slide.",
                 note="verts IS A LIST OF POSITIONS. faces IS A LIST OF INDICES INTO IT.")

# 8 ── ISSUED — P3a
s = issued_slide(d, "P3a",
             blurb="Find a simple panel in a built precedent and rebuild it from two lists — points and "
                   "faces — driven by parameters. Present it on one sheet: one studio render, one "
                   "material and a deliberate lighting setup.",
             cards=[("FIND", "A PANEL", "One precedent, simple. It has to tile 2 × 2, edge to edge."),
                    ("DECIDE", "WHAT MOVES", "Parameters first. One must change its shape, not its box."),
                    ("BUILD", "FROM TWO LISTS", "Every point calculated from the parameters. Closed mesh."),
                    ("PRESENT", "ONE RENDER", "Ground plane, one material, lighting chosen on purpose.")])
# the milestone title is long: shrink it to one line so it clears the rule
for sh in s.s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("P3a —"):
        for r in sh.text_frame.paragraphs[0].runs:
            r.font.size = Pt(21)

# 9 ── REQUIREMENTS P3a
requirements_slide(d, "P3a", sub=VISUAL)

# 10 ── NOW
now_slide(d, "NOW — BUILD THE PYRAMID",
          "Rest of the class: from a single point to a closed solid that cuts.",
          ["One point placed from an XYZ vector, and you can say why it is a position",
           "Four base corners and an apex, with Width / Depth / Height exposed",
           "All five faces joined — four sides and the base — and closure checked",
           "Cube minus pyramid working, off the same three inputs",
           "P2a submitted, and your P3a precedent panel chosen"],
          closer="LEAVE WITH A SOLID THAT CUTS. THE SHEET IS THE EASY PART.",
          bg=CYAN)

# 11 ── BEFORE NEXT CLASS
before_next_slide(d, [
    ("NEXT", "Week 5 continues descriptive geometry. We move from the pyramid to your own module, "
             "then set up its material and lighting for the P3a sheet."),
    ("FIND", "Your panel. One building, one repeating component, simple enough to write as two "
             "lists, and it has to tile 2 × 2. Credit the building, architect and source."),
    ("BUILD", "Start your module: parameters written down first, then the point list. Bring the "
              "file — next class we close it, light it and render it."),
    ("DUE", "P2a is due today. P3a, the module sheet: %s." % C.due_line("P3a")),
])

d.save(os.path.join(OUT, "ARC3133_Class04.pptx"))
print("class04 ok — %d slides" % len(d.prs.slides._sldIdLst))
