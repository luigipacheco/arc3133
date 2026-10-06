# -*- coding: utf-8 -*-
"""Class 07 — Arrays 2: conditionals (U04, Week 7).
A rule reads the index and picks one of two values for a parameter: the brick wall (k % 2),
four façade patterns for Sheet 2, and the same rules in 3D for Sheet 3. Hexagonal, radial and
sine-shaped volumes are optional practice. Studio-review week: nothing due today.

Wall images are renders of the brick-wall scene (img/class06/c05_*), coloured by position.
Pattern diagrams are drawn here, cell by cell, from the same rules shown in the code."""
import os, re
from nb import *
from slidekit import *
import coursedata as C
C.require_current_deck(7)
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img", "class06")
OUT = os.environ.get("DECK_OUT", os.path.join(HERE, ".."))
d = Deck()
WK = 7

AXT = {"i": "D7263D", "j": "3E8E00", "k": "1F6FD1"}
A32 = C.milestone("P3b")
N32, N31, N33 = A32["number"], C.milestone("P3a")["number"], C.milestone("P3c")["number"]
DUE = C.due_label("P3b")                                    # 'Sun Oct 18, 11:59 PM'
NEXT_DAY = C.date_long(WK + 1).split(" — ")[1]              # 'Oct 13'
FACADE_DAY = C.date_long(C.milestone("P3c")["release"]).split(" — ")[1]


def picture(s, name, x, y, w, h, frame=None):
    path = os.path.join(IMG, name)
    iw, ih = Image.open(path).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    if frame:
        s.rect(px, py, pw, ph, fill=None, line=frame, lw=BORDER)
    s.s.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))


TOKEN = re.compile(r"\b(i|j|k)\b")


def code(s, x, y, w, lines, size=12, step=0.3, ink=BLACK):
    """Readable Python, one line per row. i, j, k take their axis colour."""
    for ln in lines:
        runs, pos = [], 0
        body, _, comment = ln.partition("#")
        for m in TOKEN.finditer(body):
            if m.start() > pos:
                runs.append((body[pos:m.start()], "Courier New", size, None, ink))
            runs.append((m.group(1), "Courier New", size, True, AXT[m.group(1)]))
            pos = m.end()
        if pos < len(body):
            runs.append((body[pos:], "Courier New", size, None, ink))
        if comment:
            runs.append(("#" + comment, "Courier New", size, None, GREY))
        s.t(x, y, w, step, [runs or [(" ", "Courier New", size, None, ink)]])
        y += step


def grid(s, x, y, nx, ny, c, rule, on=PINK, off=CREAM):
    """Cells coloured by a rule on (i, j); j = 0 is the bottom row."""
    for j in range(ny):
        for i in range(nx):
            s.rect(x + i * c, y + (ny - 1 - j) * c, c, c, fill=on if rule(i, j) else off,
                   line=BLACK, lw=HAIR)


def sizes(s, x, y, nx, ny, c, rule, fill=LIME, turn=0, big=0.84, small=0.4):
    """The rule switches each cell's size between two values; turn rotates every square."""
    if turn:
        big, small = big * 0.72, small * 0.72
    for j in range(ny):
        for i in range(nx):
            on = rule(i, j)
            w = c * (big if on else small)
            sh = s.rect(x + i * c + (c - w) / 2, y + (ny - 1 - j) * c + (c - w) / 2, w, w,
                        fill=fill if on else CREAM, line=BLACK, lw=HAIR)
            if turn:
                sh.rotation = turn


def weights(mid):
    out = []
    for line in C.milestone(mid)["brief"]["evaluation"]:
        out.append(line.split(":**")[0].strip("* ").replace(" - ", " "))
    return " · ".join(out) + "."


# 1 ── TITLE
title_slide(d, WK, ["ARRAYS 2", "IF · SWITCH · PATTERN"],
            "A RULE THAT READS THE INDEX", bg=PINK, lfill=LIME, rfill=CYAN,
            right_tag="NOTHING DUE TODAY",
            foot="1 · 2.1 · 2.2 · 3.1 · 3.2 DUE ON TEAMS %s" % DUE.upper())

# 2 ── LAST WEEK
s = d.slide(CREAM)
s.header("LAST WEEK: EVERY COPY THE SAME", "Line, grid, cube, curve — the loop placed one module at every position.")
s.panel(L, 1.95, 5.6, 3.5, "THE COUNTERS YOU ALREADY HAVE", headfill=CYAN, tsize=14.5)
code(s, L + 0.35, 2.75, 5.0, ["i   →  along the row",
                              "j   →  which row",
                              "k   →  which layer",
                              "",
                              "index = i + j*nx + k*nx*ny"], size=13.5, step=0.38)
s.panel(6.6, 1.95, 6.05, 3.5, "TODAY", headfill=PINK)
s.t(6.9, 2.72, 5.45, 2.6, [
    s.Cb("A rule reads those counters and decides something about each copy.", 14),
    s.Cb(" ", 7),
    s.Cb("Same loop, same positions. The only new thing is a question asked of every copy — "
         "and an answer that changes one of its parameters.", 14)], ls=1.35)
s.banner(5.75, "THE INDEX WAS A COUNTER. TODAY IT BECOMES AN INPUT.", fill=YELLOW, h=0.72,
         align=PP_ALIGN.CENTER)

# 3 ── CONDITION → SWITCH → VALUE
s = d.slide(BLACK)
s.header("A CONDITION, A SWITCH, A VALUE", "Three nodes. In the class file: Compare → Switch → the input it drives.")
cards = [("CONDITION", CYAN, "k % 2 == 0", "Compare asks a yes/no question of this copy's index."),
         ("SWITCH", YELLOW, "true → A\nfalse → B", "Switch passes one of its two inputs, chosen by the answer."),
         ("VALUE", LIME, "shift · size · turn\ndepth · keep", "The chosen value drives one parameter of that copy.")]
x, wd = L, 3.55
for n, (nm, col, snippet, body) in enumerate(cards):
    s.chip(x, 1.95, wd, 0.5, nm, fill=col, size=12.5)
    s.card(x, 2.55, wd, 2.45, fill=CREAM)
    code(s, x + 0.3, 2.75, wd - 0.5, snippet.split("\n"), size=14, step=0.4)
    s.t(x + 0.3, 3.75, wd - 0.6, 1.15, [s.Cb(body, 13)], ls=1.25)
    if n < 2:
        s.t(x + wd + 0.05, 3.2, 0.6, 0.6, [s.Ab("→", 30, CREAM)], align=PP_ALIGN.CENTER)
    x += wd + 0.67
s.card(L, 5.25, W, 0.72, fill=CREAM, shadow=False)
code(s, L + 0.35, 5.42, W - 0.6, ["value = A if condition else B        # the whole idea, in one line"],
     size=15, step=0.4)
s.banner(6.3, "A CONDITION IS TRUE OR FALSE. A SWITCH TURNS THAT INTO ONE OF TWO VALUES.",
         fill=PINK, h=0.62, size=12.5, align=PP_ALIGN.CENTER)

# 4 ── MODULO
s = d.slide(CREAM)
s.header("MODULO — THE TEST THAT REPEATS", "% gives the remainder. It turns a counter that keeps growing into one that cycles.")
rows = [("i", lambda i: str(i), None),
        ("i % 2", lambda i: str(i % 2), lambda i: i % 2 == 1),
        ("i % 3", lambda i: str(i % 3), None),
        ("i % 3 == 0", lambda i: "T" if i % 3 == 0 else "F", lambda i: i % 3 == 0)]
c, x0, y = 0.82, L + 2.35, 2.05
for lab, val, hi in rows:
    s.chip(L, y, 2.15, c - 0.08, lab, fill=YELLOW if lab == "i" else CYAN, size=12.5, shadow=False)
    for i in range(9):
        on = bool(hi and hi(i))
        s.rect(x0 + i * c, y, c, c - 0.08, fill=PINK if on else CREAM, line=BLACK, lw=HAIR)
        s.t(x0 + i * c, y, c, c - 0.08, [s.Mb(val(i), 15, bold=True)], align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    y += c + 0.08
s.t(L, y + 0.1, W, 0.7, [s.Cb("Pink is where the test is true: every odd copy, then every third copy. "
                              "Read the row before you build the rule.", 14)], ls=1.3)
s.banner(6.35, "ODD/EVEN IS % 2. EVERY THIRD IS % 3 == 0. EVERY PATTERN TODAY STARTS HERE.",
         fill=LIME, h=0.62, size=12.5, align=PP_ALIGN.CENTER)

# 5 ── BRICK WALL
s = d.slide(BLACK)
s.header("01 · BRICK WALL — EVERY OTHER COURSE SHIFTS", "Arrays Part 2 · Brick wall · the row from last week, stacked in courses along a curve.")
picture(s, "c05_wall_curve.png", L, 1.95, 6.6, 4.1, frame=CREAM)
s.panel(7.55, 1.95, 5.1, 4.1, "ONE RULE ON k", headfill=PINK, tsize=14.5)
code(s, 7.8, 2.7, 4.75, ["for k in range(courses):",
                         "    shift = (k % 2) * brick / 2",
                         "    for i in range(count):",
                         "        p = point_on(curve, t)",
                         "        p = p + d * shift",
                         "        place(brick, p, z=k*h)"], size=11.5, step=0.33)
s.t(7.8, 4.85, 4.6, 1.1, [s.Cb("k is the course. Odd courses slide half a brick, even ones don't. "
                               "That single rule is what makes it a bond.", 12.5)], ls=1.25)
s.banner(6.4, "THE FIRST PATTERN: ONE CONDITION ON THE LAYER INDEX, ONE VALUE IT CHANGES.",
         fill=YELLOW, h=0.62, size=12, align=PP_ALIGN.CENTER)

# 6 ── SAME RULE, NEW PATH
s = d.slide(CREAM)
s.header("SWAP THE CURVE FOR A CIRCLE", "Same scene, same rule. Only the path that supplies the positions changed.")
picture(s, "c05_wall_circle.png", L, 1.95, 6.6, 4.1, frame=BLACK)
s.panel(7.55, 1.95, 5.1, 4.1, "TWO JOBS, KEPT APART", headfill=CYAN, tsize=14.5)
s.t(7.85, 2.72, 4.5, 3.2, [
    s.Cb("The curve answers: where is copy i?", 14),
    s.Cb(" ", 6),
    s.Cb("The rule answers: what is different about this copy?", 14),
    s.Cb(" ", 6),
    s.Cb("Because they are separate, you can change either one without touching the other. "
         "That is what makes a rule reusable.", 13.5)], ls=1.3)
s.banner(6.4, "POSITION FROM THE LOOP. DIFFERENCE FROM THE RULE.",
         fill=LIME, h=0.62, size=12.5, align=PP_ALIGN.CENTER)

# 7 ── SHEET 2: FOUR FAÇADE PATTERNS
s = d.slide(CREAM)
s.header("SHEET 2 · FOUR FAÇADE PATTERNS", "Your %s module in a 2D array. Pink = the rule is true for that copy." % N31)
pats = [("ALTERNATE ROWS", "j % 2 == 1", lambda i, j: j % 2 == 1, "true → turn 90°"),
        ("CHECKERBOARD", "(i + j) % 2 == 0", lambda i, j: (i + j) % 2 == 0, "true → scale up"),
        ("DIAGONAL BANDS", "(i + j) % 3 == 0", lambda i, j: (i + j) % 3 == 0, "true → push out"),
        ("EVERY THIRD COLUMN", "i % 3 == 0", lambda i, j: i % 3 == 0, "true → remove")]
x, wd, c = L, (W - 0.3 * 3) / 4, 0.42
for n, (nm, rule, fn, effect) in enumerate(pats):
    s.chip(x, 1.95, wd, 0.46, nm, fill=ACCENTS[n % 4], size=11)
    grid(s, x + (wd - 6 * c) / 2, 2.6, 6, 5, c, fn)
    code(s, x, 4.8, wd, [rule], size=12, step=0.32)
    s.t(x, 5.15, wd, 0.35, [s.Cb(effect, 12.5, GREY)])
    x += wd + 0.3
s.banner(5.85, "CAPTION EACH ONE: THE RULE, AND WHAT IT DOES WHEN IT IS TRUE AND WHEN IT IS FALSE.",
         fill=PINK, h=0.68, size=12.5, align=PP_ALIGN.CENTER)

# 8 ── TWO VALUES OF ONE PARAMETER
s = d.slide(BLACK)
s.header("A RULE PICKS ONE OF TWO VALUES", "Keep or remove is only one case. Any parameter with two settings can be switched.")
s.card(L, 1.95, 6.0, 4.1, fill=CREAM)
sizes(s, L + 0.6, 2.25, 7, 5, 0.69, lambda i, j: (i + j) % 2 == 0)
s.panel(7.0, 1.95, 5.65, 4.1, "THE CHECKERBOARD, DRIVING SIZE", headfill=LIME, tsize=14)
code(s, 7.3, 2.72, 5.1, ["checker = (i + j) % 2 == 0",
                         "size = big if checker else small"], size=12.5, step=0.38)
s.t(7.3, 3.65, 5.1, 2.3, [
    s.Cb("Two values: large and small. The rule decides which one each copy gets.", 13.5),
    s.Cb(" ", 6),
    s.Cb("Swap size for rotation (0° or 45°), depth, an aperture — or 'there' and 'not there'. "
         "Same rule, different parameter.", 13.5)], ls=1.3)
s.banner(6.35, "TRUE → ONE VALUE. FALSE → THE OTHER. ONE RULE DRIVES ONE PARAMETER.",
         fill=YELLOW, h=0.62, size=12.5, align=PP_ALIGN.CENTER)

# 9 ── SHEET 3: RULES IN 3D
s = d.slide(CREAM)
s.header("SHEET 3 · THE SAME RULES IN 3D", "A 3D array of cubes, read one layer at a time. Two rules: size from i + j + k, turn from k.")
x, c = L, 0.52
for k in range(3):
    s.chip(x, 1.95, 4 * c + 0.4, 0.46, "LAYER  k = %d" % k, fill=[CYAN, LIME, YELLOW][k], size=11.5)
    s.card(x, 2.55, 4 * c + 0.4, 4 * c + 0.4, fill=CREAM)
    sizes(s, x + 0.2, 2.75, 4, 4, c, lambda i, j, k=k: (i + j + k) % 2 == 0, fill=PINK,
          turn=45 if k % 2 else 0)
    x += 4 * c + 0.4 + 0.3
s.panel(8.95, 1.95, 3.7, 3.08, "READ THE STACK", headfill=PINK, tsize=13.5)
s.t(9.2, 2.65, 3.25, 2.4, [
    s.Cb("k in the size rule: the checkerboard flips every layer.", 12.5),
    s.Cb(" ", 5),
    s.Cb("k in the turn rule: odd layers turn 45°.", 12.5),
    s.Cb(" ", 5),
    s.Cb("Leave k out and every layer is identical.", 12.5)], ls=1.25)
s.card(L, 5.3, W, 0.78, fill=CREAM, shadow=False)
code(s, L + 0.3, 5.38, W - 0.6, ["size = big if (i + j + k) % 2 == 0 else small",
                                "turn = 45 if k % 2 else 0"], size=12.5, step=0.3)
s.banner(6.4, "PUT k IN THE RULE AND THE PATTERN CHANGES LAYER BY LAYER, NOT JUST ACROSS.",
         fill=CYAN, h=0.62, size=12.5, align=PP_ALIGN.CENTER)

# 10 ── SHEET 3, WRITTEN OUT
s = d.slide(BLACK)
s.header("SHEET 3, WRITTEN OUT", "The cube loop from last week, with the rules added inside it.")
s.card(L, 1.95, 7.1, 3.9, fill=CREAM)
code(s, L + 0.3, 2.15, 6.7, [
    "for k in range(layers):",
    "    for j in range(rows):",
    "        for i in range(cols):",
    "            checker = (i + j + k) % 2 == 0",
    "            size = big if checker else small",
    "            turn = 45 if k % 2 else 0",
    "            place(cube, i*dx, j*dy, k*dz,",
    "                  size=size, turn=turn)"], size=12.5, step=0.38)
s.panel(8.0, 1.95, 4.65, 3.9, "WHAT SHEET 3 SHOWS", headfill=YELLOW, tsize=14)
s.t(8.3, 2.7, 4.1, 3.0, [
    s.Cb("Two patterns, each in isometric, from the same view.", 13),
    s.Cb("At least one rule reads k.", 13),
    s.Cb("At least one switches a value — not only keep or remove.", 13),
    s.Cb("Caption: the rule, the parameter, its value when true and when false.", 13)],
    ls=1.25, space=6)
s.banner(6.2, "A HEIGHT FUNCTION OF X AND Y IS OPTIONAL NOW. THE RULES ARE NOT.",
         fill=PINK, h=0.62, size=12.5, align=PP_ALIGN.CENTER)

# 11 ── REQUIREMENTS — 3.2 (verbatim from course.yml)
requirements_slide(d, "P3b", title="REQUIREMENTS — %s" % N32, sub=weights("P3b"))

# 12 ── OPTIONAL PRACTICE
s = d.slide(CREAM)
s.header("OPTIONAL PRACTICE", "In the class file, not on the sheets. Each one is a rule you already know.")
opts = [("HEXAGONAL", CYAN, "Shift every other row by half a cell — the brick wall's j % 2, in plan."),
        ("RADIAL", LIME, "Position from radius, count and angle step. The index picks the angle."),
        ("SINE VOLUME", YELLOW, "Height from a function of X and Y. Amplitude, frequency and phase, written down.")]
x, wd = L, (W - 0.3 * 2) / 3
for nm, col, body in opts:
    s.chip(x, 1.95, wd, 0.5, nm, fill=col, size=12.5)
    s.card(x, 2.55, wd, 2.3, fill=CREAM)
    s.t(x + 0.3, 2.75, wd - 0.6, 2.0, [s.Cb(body, 14)], ls=1.3)
    x += wd + 0.3
s.banner(5.4, "DO THESE AFTER SHEETS 2 AND 3 WORK — NOT INSTEAD OF THEM.",
         fill=PINK, h=0.66, size=12.5, align=PP_ALIGN.CENTER)

# 13 ── NOW
now_slide(d, "NOW — SHEETS 2 AND 3",
          "Start at a count you can check by eye.",
          ["Brick wall: k % 2 shifting every other course, on the curve, then on a circle",
           "Your module in a 2D array with one if/Switch rule working on it",
           "Four façade rules for Sheet 2, each captioned with what it changes",
           "A 3D array of cubes where a rule switches size, rotation or height",
           "At least one rule that reads the layer index k"],
          closer="A RULE YOU CAN SAY OUT LOUD IS A RULE YOU CAN CAPTION.", bg=LIME)

# 14 ── BEFORE NEXT CLASS
before_next_slide(d, [
    ("NEXT", "%s: supported practice, nothing due. Bring all three %s sheets and anything still open."
             % (NEXT_DAY, N32)),
    ("DUE", "Everything through Arrays — 1, 2.1, 2.2 (a photo of your print), 3.1 and 3.2 — on Teams by %s. "
            "Your midterm grade is the accumulated grade of all five." % DUE),
    ("WATCH", "Rhino + Grasshopper: GH list, Arrays part 1 and part 2 are on the lesson page. "
              "The Blender Arrays part 2 video is coming."),
    ("THEN", "%s: façade attractors — %s starts." % (FACADE_DAY, N33)),
])

path = os.path.join(OUT, "ARC3133_Class07.pptx")
d.save(path)
print("class07 ok — %d slides" % len(d.prs.slides._sldIdLst))
