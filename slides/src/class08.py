# -*- coding: utf-8 -*-
"""Class 08 — Attractors (U05). P3b due. P3c and P3d issued."""
import os, math
from nb import *
from slidekit import *
import coursedata as C
import diagrams as G
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt

OUT = os.environ.get("DECK_OUT", "/home/claude/out")
d = Deck()
WK = 8

VISUAL = ("Visual quality 35% · computational understanding 25% · technical execution 20% "
          "· experimentation 10% · requirements and identity 10%.")
FABRICATION = ("Visual documentation 20% · computational understanding 20% · technical execution 15% "
               "· fabrication quality 30% · experimentation 5% · requirements and identity 10%.")

# 1 ── TITLE
title_slide(d, WK, ["ATTRACTORS"],
            "P3 CONTINUES — THE RULE STOPS COMING FROM THE INDEX", bg=CYAN, lfill=PINK, rfill=LIME)

# 2 ── AN ATTRACTOR IS A NUMBER
s = d.slide(BLACK)
s.header("AN ATTRACTOR IS A NUMBER",
         "For every point: how far is it from something? Use that number to drive a property.")
s.panel(L, 2.05, 6.2, 3.3, "THE WHOLE IDEA", headfill=LIME)
s.t(L + 0.4, 2.85, 5.4, 2.2, [
    s.Mb("for each point p:", 15),
    s.Mb("    d = distance(p, a)", 15),
    s.Mb("    value = f(d)", 15),
    s.Mb(" ", 8),
    s.Mb("Everything else this week", 13.5),
    s.Mb("is a choice of f.", 13.5)], ls=1.35)
s.panel(7.15, 2.05, 5.5, 3.3, "MEASURE IT FIRST", headfill=CYAN)
s.t(7.45, 2.85, 4.9, 2.2, [
    s.Mb("d = √((px − ax)² + (py − ay)²)", 13),
    s.Cb(" ", 6),
    s.Cb("Pythagoras, once per point. The answer is in model units.", 14),
    s.Cb(" ", 6),
    s.Cb("Index told you where in the sequence. Distance tells you how far from something "
         "outside the collection.", 14)], ls=1.3)
s.banner(5.7, "LAST WEEK YOU BUILT THE POSITIONS. THIS WEEK YOU MEASURE THEM.",
         fill=YELLOW, h=0.7, align=PP_ALIGN.CENTER)

# 3 ── READ THE FIELD FIRST
s = d.slide(BLACK)
s.header("READ THE FIELD BEFORE YOU CHANGE ANYTHING",
         "Distance is a value at every point. Look at it as grey before it becomes geometry.")
seq = [("01", "MEASURE", "Every point on a flat 2D grid gets a number: how far to the attractor.", CYAN),
       ("02", "DISPLAY", "Show those numbers as grey. This is a drawing already.", LIME),
       ("03", "REMAP", "Push that range into the range your property needs.", YELLOW),
       ("04", "DRIVE", "Only now does anything change. The field was finished first.", PINK)]
x = L
for n, ttl, body, col in seq:
    s.chip(x, 2.0, 2.85, 0.5, n + "   " + ttl, fill=col, size=12)
    s.card(x, 2.6, 2.85, 1.5, fill=CREAM)
    s.t(x + 0.25, 2.6, 2.35, 1.5, [s.Cb(body, 13)], ls=1.3, anchor=MSO_ANCHOR.MIDDLE)
    x += 3.05
s.t(L, 4.45, W, 0.3, [s.A("THE FIELD, READ AS A GREY VALUE ACROSS THE PLANE   ·   NEAR → FAR", 11, MUTE)])
G.gray_ramp(s, L, 4.78, W, 0.5)
s.banner(5.62, "IF THE GREY IMAGE IS NOT ALREADY A GOOD DRAWING, THE LINES WILL NOT SAVE IT.",
         fill=CYAN, h=0.68, align=PP_ALIGN.CENTER)
s.t(L, 6.55, W, 0.4, [s.C("P3c requires the grayscale field, with a legend, beside the line drawing. "
                          "That order is the assignment.", 13.5, MUTE)])

# 4 ── REMAP AND CLAMP
s = d.slide(CREAM)
s.header("REMAP, AND SAY WHERE YOU CLAMPED",
         "Distance comes out in model units. Your property needs a specific range.")
s.panel(L, 1.95, 6.0, 3.4, "THE OPERATION", headfill=CYAN)
s.t(L + 0.35, 2.7, 5.3, 2.3, [
    s.Mb("t   = (d − dmin) / (dmax − dmin)", 13),
    s.Mb("t   = clamp(t, 0, 1)", 13),
    s.Mb("out = outmin + t × (outmax − outmin)", 13),
    s.Mb(" ", 8),
    s.Mb("d    0 … 8.4 units    ← input", 13),
    s.Mb("out  2 … 12 mm length ← output", 13)], ls=1.4)
s.panel(7.0, 1.95, 5.65, 3.4, "SUBTRACT. DIVIDE. MULTIPLY. ADD.", headfill=LIME, tsize=14)
s.t(7.3, 2.7, 5.05, 2.3, [
    s.Cb("Nothing here is new. It is the arithmetic of a transformation, applied to a distance "
         "instead of a coordinate.", 14),
    s.Cb(" ", 7),
    s.Cb("Clamping is a decision, not a safety net. Decide what happens past the end of your "
         "range and say so.", 14)], ls=1.3)
s.banner(5.65, "P3c REQUIRES INPUT AND OUTPUT RANGES, CLAMPING AND FALLOFF LABELLED ON THE SHEET.",
         fill=YELLOW, h=0.72, align=PP_ALIGN.CENTER)

# 5 ── FALLOFF
s = d.slide(CREAM)
s.header("FALLOFF IS THE FUNCTION",
         "f is your decision. It is where the graphic quality actually comes from.")
fall = [("LINEAR", CYAN, "v = 1 − t", "Even change all the way across.", lambda t: 1 - t),
        ("INVERSE", PINK, "v = 1 / (1 + k·t)", "Sharp near the attractor, flat far off.",
         lambda t: 1 / (1 + 6 * t)),
        ("SMOOTH", LIME, "eased", "Soft at both ends. Often the best-looking.",
         lambda t: 1 - (3 * t * t - 2 * t ** 3)),
        ("STEP", YELLOW, "on / off at a cut", "No gradient at all. Two populations.",
         lambda t: 1.0 if t < 0.45 else 0.0)]
x = L
for nm, col, eq, cap, fn in fall:
    s.panel(x, 1.95, 2.85, 3.5, nm, headfill=col, tsize=14.5)
    s.t(x + 0.22, 2.68, 2.4, 0.3, [s.Mb(eq, 11.5)])
    G.falloff(s, x + 0.32, 3.9, fn)
    s.rule(x + 0.28, 4.02, 2.3, lw=Pt(1.0), color=MUTE)
    s.t(x + 0.25, 4.15, 2.35, 0.3, [s.Cb("t (remapped distance) →", 10, GREY)])
    s.t(x + 0.25, 4.55, 2.35, 0.8, [s.Cb(cap, 12.5)], ls=1.25)
    x += 3.05
s.banner(5.75, "Two students with the same attractor and different falloff produce completely "
              "different drawings. This is a design decision, not a setting.", fill=PINK, h=0.72)

# 6 ── THREE KINDS
s = d.slide(CREAM)
s.header("THREE KINDS", "Same arithmetic. What changes is what you measure the distance to.")
kinds = [("ONE POINT", CYAN, "d = distance(p, a)"),
         ("NEAREST OF SEVERAL", PINK, "d = min(distance(p, ai))"),
         ("NEAREST ON A CURVE", LIME, "d = distance(p, closest(c, p))")]
x = L
for nm, col, fm in kinds:
    s.panel(x, 1.95, 3.85, 3.45, nm, headfill=col, tsize=15)
    s.t(x + 0.2, 4.75, 3.45, 0.5, [s.Mb(fm, 10.5)], anchor=MSO_ANCHOR.MIDDLE,
        align=PP_ALIGN.CENTER)
    x += 4.08
G.field_point(s, L + 0.55, 2.75, L + 1.92, 3.5)
G.field_multi(s, L + 4.63, 2.75, [(L + 5.18, 3.05), (L + 6.83, 3.95)])
G.field_curve(s, L + 8.71, 2.75, L + 8.71, 2.9, lambda t: 3.55 - 0.62 * math.sin(t * math.pi))
s.banner(5.75, "Whatever you measure to, the output is the same thing: one number per point. "
              "P3c compares all three on the same grid, with the same mapping.",
         fill=CYAN, h=0.72)

# 7 ── PSEUDOCODE
pseudocode_slide(d, "U05", size=11,
                 sub="Measure, remap, clamp, shape the falloff, then place. In that order.",
                 note="LINE 2 IS THE ONLY LINE THAT CHANGES BETWEEN THE THREE ATTRACTOR TYPES.")

# 8 ── ISSUED — P3c
issued_slide(d, "P3c",
             blurb="Take the field off the screen. Sample it on a flat grid, turn every value into a "
                   "line, and let a pen plotter draw it. Same grid, same mapping, three attractors: "
                   "the drawing shows what the attractor alone changes.",
             cards=[("FIELD", "GREY FIRST", "The grayscale field, with a legend, before any line."),
                    ("LINES", "ONE MAPPING", "Length, rotation, spacing or hatch density, driven by distance."),
                    ("PLOT", "SVG TO PAPER", "Export at plotted size, test small, then plot the sheet.")])

# 9 ── REQUIREMENTS — P3c
s = requirements_slide(d, "P3c", sub=FABRICATION)
# the fabrication rubric is long: keep it on one line above the first row
for sh in s.s.shapes:
    if sh.has_text_frame and sh.text_frame.text == FABRICATION:
        for r in sh.text_frame.paragraphs[0].runs:
            r.font.size = Pt(12)

# 10 ── ISSUED — P3d
issued_slide(d, "P3d", badge="ALSO ISSUED TODAY", badgefill=LIME,
             blurb="Pick a real site and analyze it with Mixtli, the Blender add-on the instructor "
                   "provides. Read its output as a field — the same distance-and-remap logic as P3c — "
                   "and use it to make one spatial decision.",
             cards=[("SITE", "RECORD IT", "Source, units, scale and orientation, before any analysis."),
                    ("ANALYZE", "ONE COMPARISON", "One analysis with a legend and units, and one comparison."),
                    ("DECIDE", "YOUR READING", "The add-on computes. The question and the reading are yours.")])

# 11 ── REQUIREMENTS — P3d
requirements_slide(d, "P3d", sub=VISUAL)

# 12 ── NOW
now_slide(d, "NOW — YOUR FIRST FIELD",
          "Rest of the class: measure, look, remap, and only then change anything.",
          ["A flat 2D grid, and distance to one point measured at every sample",
           "The field displayed as grey, with a legend, before anything changes",
           "Input and output ranges written down — actual numbers, not guesses",
           "Clamp and falloff chosen, and you can say why",
           "One point swapped for several, then for a curve — same mapping throughout"],
          closer="MOVE THE ATTRACTOR. IF YOU CANNOT PREDICT WHAT HAPPENS, STOP AND READ THE FIELD.",
          bg=LIME)

# 13 ── BEFORE NEXT CLASS
cp_c = C.checkpoints("P3c")[0][1]
cp_d = C.checkpoints("P3d")[0][1]
before_next_slide(d, [
    ("NEXT", "%s: plotter workflow — field to lines to SVG to test plot — and a Mixtli working "
             "review of your site." % C.date_long(9).title()),
    ("P3c", "Checkpoint next class. %s" % cp_c),
    ("P3d", "Checkpoint next class. %s" % cp_d),
    ("PLAN", "The midterm is %s. Projects 1–3 are presented, printed and plotted. Do not leave "
             "prints or plots to the last week." % C.date_long(10).title()),
])

d.save(os.path.join(OUT, "ARC3133_Class08.pptx"))
print("class08 ok — %d slides" % len(d.prs.slides._sldIdLst))
