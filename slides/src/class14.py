# -*- coding: utf-8 -*-
"""Class 14 — Final production (U10). Production, photography, booklet."""
import os, math
from nb import *
from slidekit import *
import coursedata as C
import diagrams as G
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt

OUT = os.environ.get("DECK_OUT", "/home/claude/out")
d = Deck()
WK = 14
BOOK_RUBRIC = "Substantive revision 40% · documentation 35% · visual identity 25%."


def swap_banner(s, text, fill=PINK):
    """requirements_slide ends with a checkpoint banner (3 shapes); replace it."""
    tree = s.s.shapes._spTree
    for sp in list(s.s.shapes)[-3:]:
        tree.remove(sp._element)
    s.banner(6.77, text, fill=fill, h=0.56, size=11.5)


# 1 ── TITLE
title_slide(d, WK, ["FINAL", "PRODUCTION"],
            "P4b AND BOOK — MAKE IT, PHOTOGRAPH IT, WRITE IT UP", bg=CYAN, lfill=YELLOW, rfill=PINK)

# 2 ── PRODUCTION PLAN
s = d.slide(BLACK)
s.header("A PRODUCTION PLAN",
         "Lab time is shared and limited. Write the sequence down before you book the machine.")
plan = [("01", "CONFIRM", "Sheet, thickness, fit and the cutting file.", CYAN),
        ("02", "TEST", "Coupon and a few parts. Adjust. Regenerate.", LIME),
        ("03", "CUT", "The batch, from the regenerated file.", YELLOW),
        ("04", "CHECK", "Count and inspect every part against the list.", PINK),
        ("05", "ASSEMBLE", "In label order, dry first.", CYAN),
        ("06", "RECORD", "Photograph each stage as it happens.", LIME)]
x = L
wd = (W - 5 * 0.2) / 6
for n, ttl, body, col in plan:
    s.chip(x, 2.05, wd, 0.5, n + "  " + ttl, fill=col, size=11)
    s.card(x, 2.7, wd, 1.75, fill=CREAM)
    s.t(x + 0.18, 2.7, wd - 0.36, 1.75, [s.Cb(body, 12.5)], ls=1.3, anchor=MSO_ANCHOR.MIDDLE)
    x += wd + 0.2
s.chip(L, 4.8, 1.8, 1.0, "PLAN FOR", fill=CREAM, size=13)
s.card(L + 2.05, 4.8, W - 2.05, 1.0, fill=CREAM)
s.t(L + 2.35, 4.8, W - 2.65, 1.0, [
    s.Cb("Lab availability · the time a batch takes · a spare part or two · a retest if the fit changes", 14)],
    anchor=MSO_ANCHOR.MIDDLE)
s.banner(6.15, "TEST CUTS AND FAILED PARTS COUNT AGAINST YOUR ONE SHEET. KEEP BACKUP MATERIAL IN MIND.",
         fill=YELLOW, h=0.7, align=PP_ALIGN.CENTER)

# 3 ── TEST BEFORE BATCH
s = d.slide(CREAM)
s.header("TEST BEFORE THE BATCH",
         "A small cut now is cheaper than a sheet of parts that do not fit.")
tests = [("FIT", CYAN, "Cut the coupon from the production sheet. Confirm the slot that holds."),
         ("JOINT", LIME, "Cut one registration joint: two slices and the spine, dowel or spacer."),
         ("SMALLEST", YELLOW, "Cut the smallest and thinnest part. If it breaks, change the profile."),
         ("LABEL", PINK, "Check the engraved numbers are legible and on the hidden face.")]
y = 1.95
for nm, col, body in tests:
    s.chip(L, y, 2.0, 0.78, nm, fill=col, size=14)
    s.card(L + 2.3, y, W - 2.3, 0.78, fill=CREAM)
    s.t(L + 2.6, y, W - 2.9, 0.78, [s.Cb(body, 14)], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.98
s.banner(5.95, "ANY CHANGE AFTER A TEST — GEOMETRY OR MATERIAL — MEANS YOU REGENERATE THE CUTTING FILE.",
         fill=BLACK, color=YELLOW, h=0.7, align=PP_ALIGN.CENTER)
s.t(L, 6.85, W, 0.4, [s.C("Keep every test piece. They are evidence for the process sheet.", 13.5, GREY)])

# 4 ── ASSEMBLY SEQUENCE
s = d.slide(BLACK)
s.header("THE ASSEMBLY SEQUENCE",
         "The order you sliced in is the order you build in. Draw it before you do it.")
s.panel(L, 2.05, 5.85, 3.5, "BEFORE GLUE", headfill=CYAN)
s.t(L + 0.3, 2.8, 5.25, 2.6, [
    s.Cb("—  Lay the parts out in label order", 14),
    s.Cb("—  Dry-fit the whole assembly on its registration", 14),
    s.Cb("—  Check the up and front marks on every part", 14),
    s.Cb("—  Fix a problem part now, not after it is glued", 14)], ls=1.3, space=6)
s.panel(6.8, 2.05, 5.85, 3.5, "THE ASSEMBLY DIAGRAM", headfill=PINK)
s.t(7.1, 2.8, 5.25, 2.6, [
    s.Cb("P4b asks for one on the process sheet.", 14),
    s.Cb(" ", 6),
    s.Cb("Show the registration first, then the parts arriving in order. Number them as they "
         "are labeled.", 14),
    s.Cb(" ", 6),
    s.Cb("Someone else should be able to build it from the diagram.", 14)], ls=1.3)
s.banner(5.85, "IF THE PARTS DO NOT GO TOGETHER IN LABEL ORDER, THE ORDER IS WRONG — NOT THE PARTS.",
         fill=LIME, h=0.7, align=PP_ALIGN.CENTER)

# 5 ── PHOTOGRAPHY
s = d.slide(CREAM)
s.header("PHOTOGRAPHING PHYSICAL WORK",
         "The booklet shows the object through photographs. Treat them as drawings, with your GSM.")
ph = [("LIGHT", CYAN, "One soft main light from the side. The gaps between slices need shadow to read."),
      ("BACKGROUND", LIME, "Plain, even and wider than the frame. Nothing else on the table."),
      ("SCALE", YELLOW, "One view with a hand or a ruler, so the size is never in doubt."),
      ("SEQUENCE", PINK, "The same angle at each stage: parts, test, half built, complete.")]
x = L
for nm, col, body in ph:
    s.panel(x, 1.95, 2.85, 3.1, nm, headfill=col, tsize=14)
    s.t(x + 0.25, 2.7, 2.35, 2.2, [s.Cb(body, 13.5)], ls=1.3)
    x += 3.05
s.banner(5.35, "PHOTOGRAPH THE CSG PRINT TOO. BOOK ASKS FOR PHOTOGRAPHS AND CAPTIONS OF BOTH PHYSICAL PIECES.",
         fill=BLACK, color=CYAN, h=0.7, align=PP_ALIGN.CENTER)
s.t(L, 6.3, W, 0.7, [s.C("Shoot more than you need, from the same position each time. Pick the set on the "
                         "page, not on the camera.", 13.5, GREY)], ls=1.3)

# 6 ── DOCUMENTING REVISIONS
s = d.slide(BLACK)
s.header("DOCUMENT THE REVISIONS",
         "The booklet rewards substantive revision. Show what changed, not only the final state.")
rv = [("BEFORE", CYAN, "The version that was reviewed: sheet, file or part."),
      ("CHANGE", LIME, "What you changed, in one sentence, with the parameter or step named."),
      ("WHY", YELLOW, "The critique, test or failure that caused it."),
      ("AFTER", PINK, "The revised result, shown at the same scale and angle.")]
x = L
for nm, col, body in rv:
    s.chip(x, 2.05, 2.85, 0.5, nm, fill=col, size=13)
    s.card(x, 2.7, 2.85, 1.6, fill=CREAM)
    s.t(x + 0.25, 2.7, 2.35, 1.6, [s.Cb(body, 13)], ls=1.3, anchor=MSO_ANCHOR.MIDDLE)
    x += 3.05
s.chip(L, 4.65, 1.5, 1.0, "KEEP", fill=CREAM, size=13)
s.card(L + 1.75, 4.65, W - 1.75, 1.0, fill=CREAM)
s.t(L + 2.05, 4.65, W - 2.35, 1.0, [
    s.Cb("Dated file versions · test pieces · review notes · failed parts, photographed before you "
         "throw them away", 14)], anchor=MSO_ANCHOR.MIDDLE)
s.banner(6.1, "A REVISION YOU CANNOT SHOW IS A REVISION THE REVIEW CANNOT SEE.",
         fill=YELLOW, h=0.7, align=PP_ALIGN.CENTER)

# 7 ── BOOKLET STRUCTURE
s = d.slide(CREAM)
s.header("FOUR PROJECTS, NINE ASSIGNMENTS",
         "The booklet: 17 × 11 inch pages in your graphic standards. The GSM is a chapter, and it is also every page.")
x = L
wd = (W - 3 * 0.23) / 4
for i, p in enumerate(C.DATA["projects"]):
    ms = p["milestones"]
    s.panel(x, 1.95, wd, 3.75, "%s  %s" % (p["id"], p["title"].upper()), headfill=ACCENTS[i], tsize=13)
    yy = 2.7
    for m in ms:
        s.chip(x + 0.2, yy, 0.72, 0.42, m["id"], fill=CREAM, size=10.5, shadow=False)
        s.t(x + 1.05, yy - 0.02, wd - 1.2, 0.72,
            [s.Cb(m["title"].split(" — ")[0], 11.5)], ls=1.15)
        yy += 0.74
    if p["id"] == "P1":
        s.rule(x + 0.2, yy + 0.05, wd - 0.4, lw=Pt(1.5), color=MUTE)
        s.t(x + 0.2, yy + 0.2, wd - 0.4, 1.6, [s.Cb("Applied to every page of the booklet, and also "
                                                    "submitted as a separate revised PDF.", 11.5)], ls=1.25)
    x += wd + 0.23
s.banner(6.0, "FOR EACH: THE REVISED SHEET · THE PROCESS · WHAT CHANGED AFTER CRITIQUE · PHOTOGRAPHS FOR "
              "P2b AND P4b", fill=PINK, h=0.7, align=PP_ALIGN.CENTER, size=12)

# 8 ── REQUIREMENTS — BOOK (REMINDER)
s = requirements_slide(d, "BOOK", sub=BOOK_RUBRIC, title="REMINDER — BOOK")
swap_banner(s, "DUE AT THE FINAL REVIEW · %s   /   %d%%" % (
    C.COURSE["final_window"].upper(), C.milestone("BOOK")["weight"]), fill=PINK)

# 9 ── NOW
now_slide(d, "NOW — PRODUCTION",
          "Rest of the class: tests reviewed, production planned, booklet moving.",
          ["Your fit and joint tests on the table, with what you changed",
           "The cutting file regenerated after the last test",
           "A written production plan with your lab time booked",
           "The assembly diagram drafted in label order",
           "One booklet spread revised, with a before-and-after"],
          closer="TEST, REGENERATE, THEN CUT. IN THAT ORDER.",
          bg=YELLOW, closer_color=YELLOW)

# 10 ── BEFORE NEXT CLASS
before_next_slide(d, [
    ("NEXT", "Week 15 (target " + C.date_long(15).split("— ")[1] + ") — Assembly follow-through, "
             "documentation and file checks. Review week: nothing is due."),
    ("MAKE", "Cut the batch and assemble within your lab allocation. Photograph each stage from the "
             "same position."),
    ("BRING", "The assembly as far as it has got, your photographs, and the booklet with every "
              "project in place."),
    ("CHECK", "P4b checkpoint, " + C.date_long(15).title() + " (target): " + C.checkpoints("P4b")[1][1]),
])

d.save(os.path.join(OUT, "ARC3133_Class14.pptx"))
print("class14 ok — %d slides" % len(d.prs.slides._sldIdLst))
