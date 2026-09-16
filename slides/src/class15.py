# -*- coding: utf-8 -*-
"""Class 15 — Documentation and booklet (U10, second class). Review week, nothing due."""
import os, math
from nb import *
from slidekit import *
import coursedata as C
import diagrams as G
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt

OUT = os.environ.get("DECK_OUT", "/home/claude/out")
d = Deck()
WK = 15
BOOK_RUBRIC = "Substantive revision 40% · documentation 35% · visual identity 25%."
FAB_RUBRIC = ("Visual documentation 20% · computational understanding 20% · technical execution 15% "
              "· fabrication quality 30% · experimentation 5% · requirements and identity 10%.")
FINAL = C.COURSE["final_window"]            # "December 10–16, 2026; exact review slot to be confirmed"
FINAL_DATES = FINAL.split(";")[0]


def swap_banner(s, text, fill=PINK):
    """requirements_slide ends with a checkpoint banner (3 shapes); replace it."""
    tree = s.s.shapes._spTree
    for sp in list(s.s.shapes)[-3:]:
        tree.remove(sp._element)
    s.banner(6.77, text, fill=fill, h=0.56, size=11.5)


assert not C.milestones_due(WK), "class15 says nothing is due; course.yml disagrees"

# 1 ── TITLE
title_slide(d, WK, ["DOCUMENTATION", "AND BOOKLET"],
            "P4b AND BOOK — FINISH, CHECK, PACKAGE", bg=LIME, lfill=CYAN, rfill=PINK,
            foot="REVIEW WEEK — NO GRADED SUBMISSIONS")

# 2 ── REVIEW WEEK
s = d.slide(BLACK)
s.header("REVIEW WEEK: NOTHING IS DUE",
         "This week is reserved for studio reviews. Use the class to finish and check, not to submit.")
today = [("CHECKPOINT", CYAN, "P4b — THE ASSEMBLY", C.checkpoints("P4b")[1][1]),
         ("WORK", LIME, "ASSEMBLY FOLLOW-THROUGH", "Finish cutting and assembly within the lab time you have. "
                                                   "Replace failed parts from the same sheet."),
         ("CHECK", YELLOW, "BOOKLET AND FILES", "Every project in place, every file packaged and opened "
                                                "from a clean folder.")]
x = L
for tag, col, ttl, body in today:
    s.chip(x, 2.05, 3.85, 0.5, tag, fill=col, size=13)
    s.card(x, 2.7, 3.85, 2.6, fill=CREAM)
    s.t(x + 0.28, 2.9, 3.3, 0.4, [s.Ab(ttl, 14)])
    s.t(x + 0.28, 3.45, 3.3, 1.75, [s.Cb(body, 13)], ls=1.3)
    x += 4.08
s.banner(5.75, "NEXT AND LAST: THE FINAL REVIEW, %s." % FINAL.upper(),
         fill=PINK, h=0.7, align=PP_ALIGN.CENTER, size=12)

# 3 ── CHECKLIST — BOOK
s = requirements_slide(d, "BOOK", sub=BOOK_RUBRIC, title="FINAL CHECKLIST — BOOK")
swap_banner(s, "DUE AT THE FINAL REVIEW · %s   /   %d%%" % (FINAL.upper(), C.milestone("BOOK")["weight"]))

# 4 ── CHECKLIST — P4b
s = requirements_slide(d, "P4b", sub=None, title="FINAL CHECKLIST — P4b")
wk15, txt15 = C.checkpoints("P4b")[1]
swap_banner(s, "CHECKPOINT · %s — %s" % (C.date_long(wk15), txt15))
s.t(L, 1.4, W, 0.34, [s.C(FAB_RUBRIC, 12, GREY)])

# 5 ── BOOKLET PAGE CHECK
s = d.slide(CREAM)
s.header("CHECK EVERY SPREAD",
         "Read the booklet as a reviewer would: one project at a time, then as a single document.")
s.panel(L, 1.95, 5.85, 3.7, "EACH ASSIGNMENT SHOWS", headfill=CYAN)
s.t(L + 0.3, 2.7, 5.25, 2.85, [
    s.Cb("—  The revised result, large", 14),
    s.Cb("—  A diagram of the procedure", 14),
    s.Cb("—  The parameter values that matter", 14),
    s.Cb("—  What changed after critique, and why", 14),
    s.Cb("—  Captioned photographs, for P2b and P4b", 14)], ls=1.3, space=5)
s.panel(6.8, 1.95, 5.85, 3.7, "THE WHOLE BOOKLET", headfill=YELLOW)
s.t(7.1, 2.7, 5.25, 2.85, [
    s.Cb("—  One grid, one type system, your GSM colors", 14),
    s.Cb("—  All four projects and nine assignments", 14),
    s.Cb("—  A contents list that matches the pages", 14),
    s.Cb("—  Page numbers and consistent captions", 14),
    s.Cb("—  17 × 11 inch pages, exported to one PDF", 14)], ls=1.3, space=5)
s.banner(5.95, "SUBSTANTIVE REVISION IS 40% OF THE BOOKLET GRADE. A REPRINTED SHEET IS NOT A REVISION.",
         fill=PINK, h=0.7, align=PP_ALIGN.CENTER)

# 6 ── FILE PACKAGING
s = d.slide(BLACK)
s.header("PACKAGE THE FILES",
         "Three kinds of file. The review may open any of them, on a machine that is not yours.")
pk = [("EDITABLE", CYAN, "Every project file with its graph intact, plus its dependencies: textures, "
                         "linked files, supplied groups and add-ons you used."),
      ("FABRICATION", LIME, "The files that made the physical work: the P2b print file, the P3c SVG "
                            "and the final P4b cutting file."),
      ("PDF", YELLOW, "The booklet PDF, and the revised GSM as a separate PDF.")]
x = L
for nm, col, body in pk:
    s.panel(x, 2.05, 3.85, 2.45, nm, headfill=col, tsize=15)
    s.t(x + 0.28, 2.8, 3.3, 1.6, [s.Cb(body, 13)], ls=1.3)
    x += 4.08
s.card(L, 4.8, W, 1.25, fill=CREAM)
s.t(L + 0.35, 4.92, 7.6, 1.05, [
    s.Mb("NAME_ARC3133_FINAL/", 12, bold=True),
    s.Mb("   booklet.pdf   gsm.pdf", 12),
    s.Mb("   source/   P2a  P2b  P3a  P3b  P3c  P3d  P4a  P4b", 12)], ls=1.2)
s.t(8.4, 4.8, 3.95, 1.25, [s.Cb("An example structure only. Follow the naming and upload "
                                "instructions posted on Canvas.", 12, GREY)],
    ls=1.25, anchor=MSO_ANCHOR.MIDDLE)
s.banner(6.35, "FABRICATION FILES MUST MATCH THE OBJECT. IF YOU CHANGED THE PARTS, EXPORT THE FILE AGAIN.",
         fill=PINK, h=0.7, align=PP_ALIGN.CENTER)

# 7 ── CLEAN-FOLDER TEST
s = d.slide(CREAM)
s.header("THE CLEAN-FOLDER TEST",
         "Before you submit, pretend you are the reviewer. Do this for every editable file.")
steps = [("01", CYAN, "Copy the package to a new folder, or another computer, and open each file from there."),
         ("02", LIME, "Check nothing is missing: textures, linked files, groups, fonts in the PDFs."),
         ("03", YELLOW, "Check the units and one known dimension against the physical object."),
         ("04", PINK, "Change an exposed input. Confirm it produces the alternatives your sheets claim.")]
y = 1.95
for n, col, body in steps:
    s.chip(L, y, 0.85, 0.8, n, fill=col, size=15)
    s.card(L + 1.15, y, W - 1.15, 0.8, fill=CREAM)
    s.t(L + 1.45, y, W - 1.75, 0.8, [s.Cb(body, 14)], anchor=MSO_ANCHOR.MIDDLE)
    y += 1.0
s.banner(6.0, "IF IT ONLY OPENS ON YOUR LAPTOP, IT HAS NOT BEEN SUBMITTED.",
         fill=BLACK, color=LIME, h=0.7, align=PP_ALIGN.CENTER)

# 8 ── FINAL REVIEW FORMAT
s = d.slide(BLACK)
s.header("THE FINAL REVIEW",
         FINAL_DATES + ". The exact slot is still to be confirmed; it will be announced in class and on Canvas.")
fr = [("PRESENT", CYAN, "The laser-cut assembly, and the CSG print or its review documentation, as directed."),
      ("SHOW", LIME, "The booklet and the revised GSM, with your revisions easy to find."),
      ("EXPLAIN", YELLOW, "Open an editable file. Change an input, predict the result, explain the nodes."),
      ("ANSWER", PINK, "What was lost and gained from volume to slices, and what you would change next.")]
y = 2.05
for nm, col, body in fr:
    s.chip(L, y, 2.0, 0.78, nm, fill=col, size=14)
    s.card(L + 2.3, y, W - 2.3, 0.78, fill=CREAM)
    s.t(L + 2.6, y, W - 2.9, 0.78, [s.Cb(body, 14)], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.98
s.banner(6.1, "THE REVIEW ASSESSES THE WORK AND YOUR UNDERSTANDING OF HOW IT WAS MADE.",
         fill=YELLOW, h=0.7, align=PP_ALIGN.CENTER)

# 9 ── NOW
now_slide(d, "NOW — FINISH AND CHECK",
          "Rest of the class: the P4b checkpoint, then booklet and files.",
          ["The assembly on the table, complete or with a plan for what is left",
           "Photographs taken: light, background, scale and sequence",
           "Every booklet spread checked against the BOOK list",
           "The P4b process sheet with volume, sections, layout, assembly and photographs",
           "The package opened from a clean folder"],
          closer="NOTHING IS DUE TODAY. EVERYTHING IS DUE AT THE FINAL REVIEW.",
          bg=PINK, closer_color=PINK)

# 10 ── BEFORE THE FINAL REVIEW
s = before_next_slide(d, [
    ("WHEN", "The final review: " + FINAL + ". Watch Canvas for your slot."),
    ("FINISH", "The assembly, its photographs and the P4b process sheet. Regenerate the cutting file if "
               "any part changed."),
    ("SUBMIT", "Booklet PDF, revised GSM PDF, editable files with dependencies, and fabrication files — "
               "as listed in the BOOK checklist."),
    ("BRING", "The laser-cut assembly and the CSG print or its documentation, as directed. Be ready to "
              "open a file and change an input."),
])

d.save(os.path.join(OUT, "ARC3133_Class15.pptx"))
print("class15 ok — %d slides" % len(d.prs.slides._sldIdLst))
