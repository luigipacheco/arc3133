"""ARC3133 deck kit — reproduces the existing neo-brutalist slide language."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import math

# Palette — must stay identical to _sass/_neobrutal.scss.
# One black. Four greys, each with one job.
BLACK = "000000"; CREAM = "FFFDF7"
LIME = "C6F035"; PINK = "FF4D8D"; CYAN = "00D9E1"; YELLOW = "FFE500"
GREY = "6B6B6B"      # quiet text on a light ground
MUTE = "B4B4AC"      # quiet text on a dark ground, and guide marks
LINE = "E2E2DA"      # hairline rules inside a card
PAPER = "F4F4EE"     # the one step down from CREAM

SHADOW = 0.075
BORDER = Pt(2.75)
THIN = Pt(2.0)
HAIR = Pt(1.25)

L = 0.65
W = 12.0


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, path=None):
        """New deck, or open an existing one to append/reorder slides."""
        if path:
            self.prs = Presentation(path)
            self.layout = self.prs.slide_layouts[0]
            for lay in self.prs.slide_layouts:
                if not lay.placeholders._element.findall(qn('p:sp')):
                    self.layout = lay
                    break
        else:
            self.prs = Presentation()
            self.prs.slide_width = Inches(13.3333)
            self.prs.slide_height = Inches(7.5)
            self.layout = self.prs.slide_layouts[6]
            self.prs.slide_masters[0].slide_layouts[6].name = "DEFAULT"

    def move(self, frm, to):
        """Move the slide at index `frm` to index `to` (0-based)."""
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        el = ids[frm]
        lst.remove(el)
        lst.insert(to, el)

    def slide(self, bg=CREAM):
        s = self.prs.slides.add_slide(self.layout)
        el = s._element.find(qn('p:cSld'))
        bgel = el.makeelement(qn('p:bg'), {})
        pr = el.makeelement(qn('p:bgPr'), {})
        fill = el.makeelement(qn('a:solidFill'), {})
        fill.append(el.makeelement(qn('a:srgbClr'), {'val': bg}))
        pr.append(fill)
        pr.append(el.makeelement(qn('a:effectLst'), {}))
        bgel.append(pr)
        el.insert(0, bgel)
        return S(s, bg)

    def save(self, path):
        self.prs.save(path)


class S:
    def __init__(self, slide, bg):
        self.s = slide
        self.bg = bg
        self.dark = bg == BLACK
        self.fg = CREAM if self.dark else BLACK

    # ---------- primitives ----------
    def rect(self, x, y, w, h, fill=None, line=BLACK, lw=BORDER, shape=MSO_SHAPE.RECTANGLE):
        sh = self.s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.shadow.inherit = False
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = rgb(line); sh.line.width = lw
        sh.text_frame.text = ""
        return sh

    def card(self, x, y, w, h, fill=CREAM, line=BLACK, lw=BORDER, shadow=True):
        if shadow:
            self.rect(x + SHADOW, y + SHADOW, w, h, fill=BLACK, line=BLACK, lw=Pt(0.5))
        return self.rect(x, y, w, h, fill=fill, line=line, lw=lw)

    def rule(self, x, y, w, lw=Pt(4.0), color=None):
        sh = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), 0)
        sh.shadow.inherit = False
        sh.fill.background()
        sh.line.color.rgb = rgb(color or self.fg)
        sh.line.width = lw
        return sh

    def dot(self, cx, cy, d=0.11, fill=BLACK):
        return self.rect(cx - d / 2, cy - d / 2, d, d, fill=fill, line=fill,
                         lw=Pt(0.5), shape=MSO_SHAPE.OVAL)

    def t(self, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          space=0, ls=None):
        """paras: list of paragraphs; each is a str-tuple (text,font,size,bold,color)
        or a list of such tuples."""
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        if paras and isinstance(paras[0], tuple):
            paras = [[p] for p in paras]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if space:
                p.space_after = Pt(space)
            if ls:
                p.line_spacing = ls
            for (txt, fnt, sz, bold, col) in para:
                r = p.add_run(); r.text = txt
                r.font.name = fnt; r.font.size = Pt(sz)
                r.font.bold = bold; r.font.color.rgb = rgb(col)
        return tb

    # run shorthands
    def A(self, txt, sz, col=None, bold=True):
        return (txt, "Arial", sz, bold, col or self.fg)

    def C(self, txt, sz, col=None, bold=None):
        return (txt, "Calibri", sz, bold, col or self.fg)

    def M(self, txt, sz, col=None, bold=None):
        return (txt, "Courier New", sz, bold, col or self.fg)

    # black-on-light runs for text sitting on coloured cards
    def Ab(self, txt, sz, col=BLACK, bold=True):
        return (txt, "Arial", sz, bold, col)

    def Cb(self, txt, sz, col=BLACK, bold=None):
        return (txt, "Calibri", sz, bold, col)

    def Mb(self, txt, sz, col=BLACK, bold=None):
        return (txt, "Courier New", sz, bold, col)

    # ---------- compound blocks ----------
    def header(self, title, sub=None, tsize=33):
        self.rule(L, 0.62, W)
        self.t(L, 0.74, W, 0.62, [self.A(title, tsize)])
        if sub:
            self.t(L, 1.38, W, 0.34, [self.C(sub, 13.5, MUTE if self.dark else GREY)])

    def chip(self, x, y, w, h, label, fill=YELLOW, size=11, shadow=True):
        self.card(x, y, w, h, fill=fill, shadow=shadow)
        self.t(x, y + (h - 0.22) / 2, w, 0.3, [self.Ab(label, size)], align=PP_ALIGN.CENTER)

    def banner(self, y, text, fill=PINK, color=BLACK, h=0.68, size=12.5,
               x=L, w=W, align=PP_ALIGN.LEFT, font="Arial", bold=True):
        self.card(x, y, w, h, fill=fill)
        pad = 0.35 if align == PP_ALIGN.LEFT else 0.2
        self.t(x + pad, y, w - 2 * pad, h, [(text, font, size, bold, color)],
               align=align, anchor=MSO_ANCHOR.MIDDLE)

    def numrow(self, y, chip, body, chipfill=LIME, cw=1.5, h=0.5, csize=12.5, size=14):
        self.chip(L, y, cw, h, chip, fill=chipfill, size=csize)
        self.t(L + cw + 0.35, y, W - cw - 0.35, h, [self.Cb(body, size)],
               anchor=MSO_ANCHOR.MIDDLE)

    def checkrow(self, y, body, h=0.62, mark="✓"):
        self.chip(L, y, 0.68, h, mark, fill=CREAM, size=17)
        self.card(L + 0.98, y, W - 0.98, h, fill=CREAM)
        self.t(L + 1.28, y, W - 1.58, h, [self.Cb(body, 14.5)], anchor=MSO_ANCHOR.MIDDLE)

    def actionrow(self, y, label, body, fill=CYAN, h=0.88):
        self.chip(L, y, 1.65, h, label, fill=fill, size=14)
        self.card(L + 1.95, y, W - 1.95, h, fill=CREAM)
        self.t(L + 2.25, y + 0.13, W - 2.55, h - 0.26, [self.Cb(body, 13.5)], ls=1.2)

    def panel(self, x, y, w, h, title, headfill=LIME, fill=CREAM, headh=0.55, tsize=15.875):
        self.card(x, y, w, h, fill=fill)
        if headfill:
            self.rect(x, y, w, headh, fill=headfill, line=BLACK, lw=BORDER)
        self.t(x + 0.25, y, w - 0.5, headh, [self.Ab(title, tsize)], anchor=MSO_ANCHOR.MIDDLE)

    def title_slide(self, kicker, big_lines, left_tag, right_tag, foot,
                    lfill=PINK, rfill=CYAN):
        self.rect(L, 1.15, W, 0.28, fill=BLACK, line=None)
        self.t(L, 1.62, 8.0, 0.5, [self.Ab(kicker, 18)])
        self.card(L, 2.25, 11.1, 2.35, fill=CREAM)
        self.t(1.1, 2.5, 10.2, 1.9, [self.Ab(x, 40) for x in big_lines],
               anchor=MSO_ANCHOR.MIDDLE, ls=1.15)
        self.card(L, 4.9, 7.4, 0.62, fill=lfill)
        self.card(8.45, 4.9, 3.3, 0.62, fill=rfill)
        self.t(0.95, 4.9, 7.0, 0.62, [self.Ab(left_tag, 13)], anchor=MSO_ANCHOR.MIDDLE)
        self.t(8.45, 4.9, 3.3, 0.62, [self.Ab(right_tag, 13)],
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self.t(L, 6.42, W, 0.32, [self.Ab(foot, 10.5)])

    def issued(self, badge, title, meta, blurb, badgefill=YELLOW):
        self.chip(L, 0.62, 2.4, 0.42, badge, fill=badgefill, size=11)
        self.t(L, 1.26, W, 0.62, [self.A(title, 32, CREAM)])
        self.rule(L, 1.98, W, lw=Pt(4.0), color=CREAM)
        self.t(L, 2.11, W, 0.36, [self.A(meta, 13, YELLOW)])
        self.t(L, 2.6, 11.6, 1.0, [self.C(blurb, 15, CREAM)], ls=1.3)
